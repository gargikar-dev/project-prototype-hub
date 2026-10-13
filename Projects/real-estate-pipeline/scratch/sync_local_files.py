import os
import sys
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.path.append(r"c:\Users\crypt\OneDrive\Desktop\Real-Estate-Lead-Gen")
sys.path.append(r"c:\Users\crypt\OneDrive\Desktop\Real-Estate-Lead-Gen\scratch")

# Load previous 16 leads from scrape_linkedin_leads.py
from scrape_linkedin_leads import LEADS_ICP_48H as BATCH_1
from prepare_new_leads import NEW_15_LEADS as BATCH_2


# Normalize batch 2 keys to match batch 1
batch_2_norm = []
for b in BATCH_2:
    batch_2_norm.append({
        "date_posted": b["date_posted"],
        "company_name": b["company_name"],
        "role_needed": b["role_needed"],
        "linkedin_url": b["linkedin_url"],
        "city_state": b["city_state"],
        "country": b["country"],
        "company_size": b["company_size"],
        "email_contact": b["email_contact"],
        "deal_value": b["deal_value"],
        "remarks": b["remarks"]
    })

ALL_31_LEADS = BATCH_1 + batch_2_norm

workspace_dir = r"c:\Users\crypt\OneDrive\Desktop\Real-Estate-Lead-Gen"
csv_file = os.path.join(workspace_dir, "leads_linkedin_icp_us_uk_48h.csv")
xlsx_file = os.path.join(workspace_dir, "leads_linkedin_icp_us_uk_48h.xlsx")
master_xlsx = os.path.join(workspace_dir, "leads_linkedin_icp.xlsx")
md_file = os.path.join(workspace_dir, "leads_linkedin_icp_us_uk_48h.md")

# 1. Update CSV
with open(csv_file, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow([
        "Date Posted (<48h)", "Company Name", "Role / Lead Need", "Verified LinkedIn URL",
        "City / State", "Country", "Company Size", "Contact Email", "Deal / Contract Value", "Remarks & ICP Fit"
    ])
    for lead in ALL_31_LEADS:
        writer.writerow([
            lead["date_posted"], lead["company_name"], lead["role_needed"], lead["linkedin_url"],
            lead["city_state"], lead["country"], lead["company_size"], lead["email_contact"],
            lead["deal_value"], lead["remarks"]
        ])

print(f"Updated CSV with {len(ALL_31_LEADS)} leads: {csv_file}")

# 2. Update Excel
def build_excel(target_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master Leads"

    ws.merge_cells("A1:K1")
    t_cell = ws["A1"]
    t_cell.value = "Master ICP Lead Generation Directory (< 48 Hours Old | US & UK)"
    t_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    headers = [
        "#", "Date Posted (<48h)", "Company Name", "Role / Lead Need", "Verified LinkedIn URL",
        "City / State", "Country", "Company Size", "Contact Email", "Deal / Contract Value", "Remarks & ICP Fit"
    ]
    ws.append([])
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    row_white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    link_font = Font(name="Calibri", size=10, color="0284C7", underline="single")
    text_font = Font(name="Calibri", size=10, color="0F172A")
    bold_text_font = Font(name="Calibri", size=10, bold=True, color="0F172A")

    current_row = 4
    for idx, lead in enumerate(ALL_31_LEADS, start=1):
        fill = row_alt_fill if idx % 2 == 0 else row_white_fill
        row_values = [
            idx, lead["date_posted"], lead["company_name"], lead["role_needed"],
            lead["linkedin_url"], lead["city_state"], lead["country"], lead["company_size"],
            lead["email_contact"], lead["deal_value"], lead["remarks"]
        ]
        ws.append(row_values)
        ws.row_dimensions[current_row].height = 36

        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = fill
            cell.border = thin_border
            cell.font = text_font

            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_text_font
            elif col_idx == 3:
                cell.font = bold_text_font
                cell.alignment = Alignment(vertical="center")
            elif col_idx == 5:
                cell.font = link_font
                cell.hyperlink = lead["linkedin_url"]
                cell.value = "Direct LinkedIn Posting"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (2, 6, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        current_row += 1

    wb.save(target_path)
    print(f"Saved: {target_path}")

build_excel(xlsx_file)
build_excel(master_xlsx)

# 3. Update Markdown document
with open(md_file, "w", encoding="utf-8") as f:
    f.write("# Master ICP Lead Generation Directory (< 48 Hours Old | US & UK)\n\n")
    f.write(f"**Total Qualified Leads:** {len(ALL_31_LEADS)}\n\n")
    f.write("- **Posting Velocity:** Strictly < 48 hours old (posted between 4 mins and 24 hours ago)\n")
    f.write("- **Geography:** United States (US) & United Kingdom (UK)\n")
    f.write("- **Company Size:** 10 – 30 employees\n")
    f.write("- **Deal Value:** $7,000+\n")
    f.write("- **Single Master Google Sheet:** [Lead Generation ICP (Tab: Master Leads)](https://docs.google.com/spreadsheets/d/1k8aPzpsR3E-yIGsjEpr87cGrzwhgHqUjw-SSagyxSiQ/edit)\n\n")
    f.write("### 📋 15 Newly Appended Leads (Batch 2)\n\n")
    f.write("| # | Date Posted (<48h) | Company Name | Role / Lead Need | Location | Country | Co. Size | Deal / Value | Live Link | Remarks |\n")
    f.write("| :-: | :--- | :--- | :--- | :--- | :--- | :--- | :--- | :--- | :--- |\n")
    for idx, l in enumerate(batch_2_norm, start=17):
        f.write(f"| **{idx}** | {l['date_posted']} | **{l['company_name']}** | {l['role_needed']} | {l['city_state']} | {l['country']} | {l['company_size']} | {l['deal_value']} | [View Posting]({l['linkedin_url']}) | {l['remarks']} |\n")

print(f"Updated Markdown: {md_file}")
