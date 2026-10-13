#!/usr/bin/env python3
"""
Sheet Exporter for Reddit ICP Lead Engine
-----------------------------------------
Manages synchronization and persistence of leads across:
1. Dedicated Excel Tracker: reports/reddit_leads_tracker.xlsx
2. Dedicated CSV Tracker: reports/reddit_leads_tracker.csv
3. Dedicated Google Sheet: Formats and prepares records for Google Sheets sync.

Ensures deduplication across multiple executions:
Running the pipeline repeatedly appends fresh leads without duplicating existing posts.
"""

import os
import csv
from datetime import datetime
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Timestamp", 20),
    ("Subreddit", 16),
    ("Post ID", 14),
    ("Thread Title", 45),
    ("Thread URL", 45),
    ("Author", 18),
    ("Composite Score", 16),
    ("Intent Score", 14),
    ("ICP Match Score", 16),
    ("Intent Category", 20),
    ("Target / Competitor", 25),
    ("Core Frustration / Summary", 50),
    ("Recommended Angle", 45),
    ("Staged 80/20 Response Draft", 60),
    ("Spam Risk Score", 16),
    ("Compliance Status", 18),
    ("HITL Status", 16),
]

HEADER_FILL = PatternFill(start_color="1A73E8", end_color="1A73E8", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=10)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

def lead_to_row(lead: dict[str, Any]) -> list[Any]:
    """Convert a lead dictionary into a flat row matching the sheet columns."""
    marketing = lead.get("marketing_engagement", {})
    risk = marketing.get("risk_assessment", {})
    compliance = "PASS" if risk.get("rule_compliance", True) else "FLAGGED"
    
    return [
        datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"),
        f"r/{lead.get('subreddit', '')}".replace("r/r/", "r/"),
        lead.get("post_id", ""),
        lead.get("thread_title", ""),
        lead.get("url", ""),
        f"u/{lead.get('author', 'anonymous')}".replace("u/u/", "u/"),
        lead.get("composite_score", 0),
        lead.get("intent_score", 0),
        lead.get("icp_match_score", 0),
        lead.get("intent_category", ""),
        lead.get("target_product", ""),
        lead.get("core_problem_summary", ""),
        lead.get("suggested_angle", ""),
        marketing.get("drafted_response", ""),
        risk.get("spam_risk_score", 0.05),
        compliance,
        "Pending Review" if marketing.get("requires_hitl_approval", True) else "Approved",
    ]

def append_to_excel(leads: list[dict[str, Any]], excel_path: str = "reports/reddit_leads_tracker.xlsx") -> int:
    """
    Append new qualified leads to the dedicated Excel spreadsheet.
    Deduplicates by Post ID and URL so running multiple times keeps historical data clean.
    Returns the count of newly appended leads.
    """
    os.makedirs(os.path.dirname(os.path.abspath(excel_path)), exist_ok=True)
    
    if os.path.exists(excel_path):
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        # Identify existing Post IDs and URLs from column C and E
        existing_keys = set()
        for r in range(2, ws.max_row + 1):
            post_id = ws.cell(row=r, column=3).value
            url = ws.cell(row=r, column=5).value
            if post_id:
                existing_keys.add(str(post_id).strip())
            if url:
                existing_keys.add(str(url).strip())
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Qualified Leads"
        existing_keys = set()
        
        # Write and style header row
        ws.row_dimensions[1].height = 28
        for col_idx, (col_name, col_width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = ALIGN_CENTER
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_width
        
        ws.freeze_panes = "A2"

    added_count = 0
    start_row = ws.max_row + 1

    for lead in leads:
        lead_id = str(lead.get("post_id", "")).strip()
        lead_url = str(lead.get("url", "")).strip()
        
        # Deduplication check
        if (lead_id and lead_id in existing_keys) or (lead_url and lead_url in existing_keys):
            continue
            
        row_values = lead_to_row(lead)
        current_row = start_row + added_count
        ws.row_dimensions[current_row].height = 40
        
        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = ROW_FONT
            if col_idx in [1, 2, 3, 6, 7, 8, 9, 10, 15, 16, 17]:
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT
                
        if lead_id:
            existing_keys.add(lead_id)
        if lead_url:
            existing_keys.add(lead_url)
        added_count += 1

    wb.save(excel_path)
    return added_count

def append_to_csv(leads: list[dict[str, Any]], csv_path: str = "reports/reddit_leads_tracker.csv") -> int:
    """Append leads to CSV with deduplication."""
    os.makedirs(os.path.dirname(os.path.abspath(csv_path)), exist_ok=True)
    existing_keys = set()
    file_exists = os.path.exists(csv_path)

    if file_exists:
        try:
            with open(csv_path, "r", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                header = next(reader, None)
                for row in reader:
                    if len(row) >= 5:
                        existing_keys.add(row[2].strip())
                        existing_keys.add(row[4].strip())
        except Exception:
            pass

    headers = [col[0] for col in COLUMNS]
    added_count = 0

    with open(csv_path, "a", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(headers)

        for lead in leads:
            lead_id = str(lead.get("post_id", "")).strip()
            lead_url = str(lead.get("url", "")).strip()
            if (lead_id and lead_id in existing_keys) or (lead_url and lead_url in existing_keys):
                continue
            writer.writerow(lead_to_row(lead))
            if lead_id:
                existing_keys.add(lead_id)
            if lead_url:
                existing_keys.add(lead_url)
            added_count += 1

    return added_count
