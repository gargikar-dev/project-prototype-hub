#!/usr/bin/env python3
import sys, os, json, glob
sys.path.append('scripts')
sys.stdout.reconfigure(encoding='utf-8')
from ai_lead_scorer import evaluate_lead
from marketing_drafter import draft_marketing_response
from sheet_exporter import lead_to_row, COLUMNS, HEADER_FILL, HEADER_FONT, ROW_FONT, ALIGN_CENTER, ALIGN_LEFT
import openpyxl
from openpyxl.utils import get_column_letter

icp = json.load(open('config/ai_business_icp.json', encoding='utf-8'))
product_profile = icp.get('product_offering', {})

seen = set()
all_candidates = []

for f in glob.glob('.cache/*.json'):
    try:
        posts = json.load(open(f, encoding='utf-8'))
        for p in posts:
            p_id = p.get('id') or p.get('url')
            if not p_id or p_id in seen:
                continue
            seen.add(p_id)
            evaluated = evaluate_lead(p, icp)
            if not evaluated.get('is_disqualified') and evaluated.get('composite_score', 0) >= 71:
                all_candidates.append(evaluated)
    except Exception as e:
        pass

all_candidates.sort(key=lambda x: x.get('composite_score', 0), reverse=True)
print(f"Total pristine qualified candidates (Score > 70 & < 24h & Whitelisted): {len(all_candidates)}")

# Take the top 20 pristine leads
top_20 = all_candidates[:20]
print(f"Selected top {len(top_20)} leads for Excel Tracker.")

# Add marketing drafts
staged_leads = []
for lead in top_20:
    draft = draft_marketing_response(lead, product_profile)
    staged_leads.append({**lead, "marketing_engagement": draft})

# Rebuild Excel tracker cleanly
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Verified ICP Leads"
ws.row_dimensions[1].height = 28

for col_idx, (col_name, col_width) in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=col_idx, value=col_name)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = ALIGN_CENTER
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = col_width

ws.freeze_panes = "A2"

for row_idx, lead in enumerate(staged_leads, start=2):
    ws.row_dimensions[row_idx].height = 42
    row_vals = lead_to_row(lead)
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = ROW_FONT
        if col_idx in [1, 2, 3, 6, 7, 8, 9, 10, 15, 16, 17]:
            cell.alignment = ALIGN_CENTER
        else:
            cell.alignment = ALIGN_LEFT

excel_path = "reports/reddit_leads_tracker.xlsx"
wb.save(excel_path)
print(f"[OK] Cleanly rebuilt {excel_path} with {len(staged_leads)} verified leads.")

# Rebuild CSV tracker cleanly
import csv
csv_path = "reports/reddit_leads_tracker.csv"
with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
    writer = csv.writer(f)
    writer.writerow([c[0] for c in COLUMNS])
    for lead in staged_leads:
        writer.writerow(lead_to_row(lead))
print(f"[OK] Cleanly rebuilt {csv_path} with {len(staged_leads)} verified leads.")

# Update latest_leads.json and latest_leads.md
from run_pipeline import generate_markdown_report
summary_payload = {
    "timestamp": "2026-10-10T06:55:00",
    "icp_name": icp.get("icp_name"),
    "total_harvested": len(seen),
    "total_qualified": len(staged_leads),
    "min_score_threshold": 71,
    "max_age_hours": 24,
    "dedicated_google_sheet": "https://docs.google.com/spreadsheets/d/1BKv_Uh3FynfdmUYX4o4yvuHdV8B6yxyYg5LURg4rHyg/edit",
    "leads": staged_leads
}

with open("reports/latest_leads.json", "w", encoding="utf-8") as f:
    json.dump(summary_payload, f, indent=2)

generate_markdown_report(summary_payload, "reports/latest_leads.md")
print(f"[OK] Updated latest_leads.json and latest_leads.md")

for i, l in enumerate(staged_leads, 1):
    print(f"{i}. [{l['composite_score']}] r/{l['subreddit']} | u/{l['author']} : {l['thread_title'][:60]}")
